import os
from pptx import Presentation
import pymupdf

def extract_pptx_text(file_path: str) -> list[dict]:
    """Extracts text from a PowerPoint file, preserving slide numbers as traceability markers."""
    chunks = []
    prs = Presentation(file_path)
    file_name = os.path.basename(file_path)
    
    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_text.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        slide_text.append(" | ".join(row_text))
                        
        if slide_text:
            chunks.append({
                "file_name": file_name,
                "location_marker": f"Slide {slide_idx}",
                "content": "\n".join(slide_text)
            })
    return chunks

def extract_pdf_text(file_path: str) -> list[dict]:
    """Extracts text from a PDF, automatically falling back to OCR for scanned image pages."""
    chunks = []
    file_name = os.path.basename(file_path)
    
    with pymupdf.open(file_path) as doc:
        print(f"[*] Processing PDF: {file_name} ({len(doc)} pages total)")
        for page_idx, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            
            # If standard text extraction yields nothing, use PyMuPDF's built-in OCR
            if not text:
                try:
                    tp = page.get_textpage_ocr(language="eng")
                    text = page.get_text(textpage=tp).strip()
                except Exception as e:
                    print(f"[!] OCR execution warning on Page {page_idx}: {e}")
            
            if text:
                chunks.append({
                    "file_name": file_name,
                    "location_marker": f"Page {page_idx}",
                    "content": text
                })
            else:
                print(f"[!] Warning: Page {page_idx} yielded no text.")
                
    return chunks

def extract_document(file_path: str) -> list[dict]:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pptx":
        return extract_pptx_text(file_path)
    elif ext == ".pdf":
        return extract_pdf_text(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")