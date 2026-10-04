import os
from abc import ABC, abstractmethod
import docx
from PIL import Image
from pypdf import PdfReader


class BaseDocumentHandler(ABC):
    """Abstract Base Class for all document handlers."""

    def __init__(self, file_path: str):
        self.file_path = file_path
        self.filename = os.path.basename(file_path)

    @abstractmethod
    def get_file_type(self) -> str:
        """Return human-readable file category name."""
        pass

    @abstractmethod
    def get_page_count(self) -> int:
        """Return total number of pages."""
        pass

    @abstractmethod
    def get_pages_text(self) -> list[str]:
        """Return list of text strings partitioned per page."""
        pass


class TextDocumentHandler(BaseDocumentHandler):

    def get_file_type(self) -> str:
        return "Text Document (.txt)"

    def get_page_count(self) -> int:
        return 1

    def get_pages_text(self) -> list[str]:
        with open(self.file_path, "r", encoding="utf-8", errors="ignore") as f:
            return [f.read()]


class PDFDocumentHandler(BaseDocumentHandler):

    def get_file_type(self) -> str:
        return "PDF Document (.pdf)"

    def get_page_count(self) -> int:
        try:
            reader = PdfReader(self.file_path)
            return len(reader.pages)
        except Exception:
            return 0

    def get_pages_text(self) -> list[str]:
        try:
            reader = PdfReader(self.file_path)
            return [
                page.extract_text() or "[No selectable text on page]"
                for page in reader.pages
            ]
        except Exception as e:
            return [f"Error reading PDF: {str(e)}"]


class WordDocumentHandler(BaseDocumentHandler):

    def get_file_type(self) -> str:
        return "Word Document (.docx)"

    def get_page_count(self) -> int:
        try:
            doc = docx.Document(self.file_path)
            page_breaks = 0
            for p in doc.paragraphs:
                for run in p.runs:
                    if (
                        "lastRenderedPageBreak" in run._element.xml
                        or 'w:type="page"' in run._element.xml
                    ):
                        page_breaks += 1
            return max(1, page_breaks + 1)
        except Exception:
            return 1

    def get_pages_text(self) -> list[str]:
        try:
            doc = docx.Document(self.file_path)
            full_text = "\n\n".join(
                [p.text for p in doc.paragraphs if p.text.strip()]
            )
            return [full_text]
        except Exception as e:
            return [f"Error reading Word document: {str(e)}"]


class ImageDocumentHandler(BaseDocumentHandler):

    def get_file_type(self) -> str:
        ext = os.path.splitext(self.filename)[1].upper()
        return f"Image File ({ext})"

    def get_page_count(self) -> int:
        return 1

    def get_pages_text(self) -> list[str]:
        try:
            with Image.open(self.file_path) as img:
                return [
                    f"[Image File Uploaded]\nFormat: {img.format}\nDimensions: {img.width} x {img.height} px\nColor Mode: {img.mode}"
                ]
        except Exception as e:
            return [f"Error inspecting image: {str(e)}"]


class DocumentHandlerFactory:
    """Factory class to detect and return the appropriate Handler instance."""

    @staticmethod
    def get_handler(file_path: str) -> BaseDocumentHandler:
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".txt":
            return TextDocumentHandler(file_path)
        elif ext == ".pdf":
            return PDFDocumentHandler(file_path)
        elif ext in [".docx", ".doc"]:
            return WordDocumentHandler(file_path)
        elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"]:
            return ImageDocumentHandler(file_path)
        else:
            raise ValueError(f"Unsupported file format: '{ext}'")